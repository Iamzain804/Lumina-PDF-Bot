"""V2 Vector Store — semantic embeddings + FAISS + hybrid search.

Drop-in replacement for DocumentVectorStore. Same public interface,
completely different internals. Safe to run alongside v1.
"""

from pathlib import Path
from typing import List, Tuple

import numpy as np
import pdfplumber
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from rank_bm25 import BM25Okapi

try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class DocumentVectorStoreV2:
    """V2: semantic embeddings, FAISS index, hybrid BM25+dense search."""

    def __init__(self, config):
        self.config = config
        self.embeddings = HuggingFaceEmbeddings(
            model_name=config.EMBEDDING_MODEL,
            model_kwargs={"device": "cpu"},
            encode_kwargs={"normalize_embeddings": True},
        )
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=config.CHUNK_SIZE,
            chunk_overlap=config.CHUNK_OVERLAP,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    # ------------------------------------------------------------------ #
    #  Document Loading                                                    #
    # ------------------------------------------------------------------ #

    def load_document(self, file_path: str) -> str:
        """Extract text — kept for RAGEngine.get_document_summary compat."""
        path = Path(file_path)
        ext = path.suffix.lower()
        if ext == ".pdf":
            return self._load_pdf_text(path)
        elif ext in (".txt", ".md"):
            return self._load_text(path)
        elif ext == ".docx":
            return self._load_docx(path)
        raise ValueError(f"Unsupported format: {ext}")

    def load_document_with_pages(self, file_path: str) -> List[Document]:
        """Load document as LangChain Documents with page metadata."""
        path = Path(file_path)
        ext = path.suffix.lower()
        if ext == ".pdf":
            return self._load_pdf_pages(path)
        else:
            text = self.load_document(file_path)
            return [Document(page_content=text, metadata={"page": 1, "source": path.name})]

    def _load_pdf_pages(self, path: Path) -> List[Document]:
        docs = []
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append(Document(
                        page_content=text,
                        metadata={"page": i, "source": path.name}
                    ))
        if not docs:
            raise ValueError("No text extracted from PDF")
        return docs

    def _load_pdf_text(self, path: Path) -> str:
        with pdfplumber.open(path) as pdf:
            text = "\n".join(p.extract_text() or "" for p in pdf.pages)
        if not text.strip():
            raise ValueError("No text found in PDF")
        return text

    def _load_text(self, path: Path) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                text = path.read_text(encoding=enc)
                if text.strip():
                    return text
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not decode text file")

    def _load_docx(self, path: Path) -> str:
        if not DOCX_AVAILABLE:
            raise ValueError("python-docx not installed")
        doc = DocxDocument(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        text = "\n".join(parts)
        if not text.strip():
            raise ValueError("No text found in DOCX")
        return text

    # ------------------------------------------------------------------ #
    #  Chunking                                                            #
    # ------------------------------------------------------------------ #

    def split_text(self, text: str) -> List[str]:
        """Compat method — RAGEngine calls this directly."""
        chunks = self.text_splitter.split_text(text)
        if not chunks:
            raise ValueError("No chunks created")
        return chunks

    def split_documents(self, docs: List[Document]) -> List[Document]:
        """Split with page metadata preserved per chunk."""
        chunks = []
        for doc in docs:
            splits = self.text_splitter.split_text(doc.page_content)
            for i, chunk in enumerate(splits):
                chunks.append(Document(
                    page_content=chunk,
                    metadata={**doc.metadata, "chunk_index": i}
                ))
        return chunks

    # ------------------------------------------------------------------ #
    #  Embeddings + FAISS                                                  #
    # ------------------------------------------------------------------ #

    def create_embeddings(self, chunks: List[str]):
        """Create FAISS vectorstore from plain text chunks."""
        docs = [Document(page_content=c, metadata={"chunk_index": i})
                for i, c in enumerate(chunks)]
        return FAISS.from_documents(docs, self.embeddings)

    def create_embeddings_from_docs(self, chunks: List[Document]):
        """Create FAISS vectorstore from Document objects (with metadata)."""
        return FAISS.from_documents(chunks, self.embeddings)

    def save_vectorstore(self, vectorstore, doc_name: str) -> None:
        save_path = Path(self.config.VECTOR_STORE_DIR) / doc_name
        save_path.mkdir(parents=True, exist_ok=True)
        vectorstore.save_local(str(save_path))

    def load_vectorstore(self, doc_name: str):
        load_path = Path(self.config.VECTOR_STORE_DIR) / doc_name
        if not load_path.exists():
            return None
        try:
            return FAISS.load_local(
                str(load_path),
                self.embeddings,
                allow_dangerous_deserialization=True,
            )
        except Exception:
            return None

    # ------------------------------------------------------------------ #
    #  Search                                                              #
    # ------------------------------------------------------------------ #

    def search(self, query: str, vectorstore, k: int) -> List[Tuple]:
        """Hybrid search: dense (FAISS) + sparse (BM25) → RRF merge."""
        # 1. Dense retrieval
        dense_results = vectorstore.similarity_search_with_score(query, k=k * 3)

        # 2. BM25 sparse retrieval on same candidate pool
        all_docs = [doc for doc, _ in dense_results]
        tokenized = [doc.page_content.lower().split() for doc in all_docs]

        if len(tokenized) > 1:
            bm25 = BM25Okapi(tokenized)
            bm25_scores = bm25.get_scores(query.lower().split())
        else:
            bm25_scores = np.ones(len(all_docs))

        # 3. Reciprocal Rank Fusion
        dense_rank = {doc.page_content: rank for rank, (doc, _) in enumerate(dense_results)}
        bm25_rank  = {all_docs[i].page_content: rank
                      for rank, i in enumerate(np.argsort(bm25_scores)[::-1])}

        rrf_scores = {}
        K = 60  # RRF constant
        for content in dense_rank:
            d_r = dense_rank.get(content, len(all_docs))
            b_r = bm25_rank.get(content, len(all_docs))
            rrf_scores[content] = 1 / (K + d_r) + 1 / (K + b_r)

        # 4. Re-sort by RRF, apply score threshold, return top-k
        threshold = getattr(self.config, "RETRIEVAL_SCORE_THRESHOLD", 0.0)
        doc_map = {doc.page_content: (doc, score) for doc, score in dense_results}

        results = []
        for content, _ in sorted(rrf_scores.items(), key=lambda x: x[1], reverse=True)[:k]:
            doc, dense_score = doc_map[content]
            normalized = float(1 - dense_score / 2) if dense_score <= 2 else 0.0
            if normalized >= threshold:
                results.append((doc, normalized))

        return results
